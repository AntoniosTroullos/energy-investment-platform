from __future__ import annotations

import importlib.util
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document


SCRIPT = Path(__file__).with_name("build_docx.py")
SPEC = importlib.util.spec_from_file_location("thesis_build_docx", SCRIPT)
assert SPEC is not None and SPEC.loader is not None
build_docx = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(build_docx)


class NativeEquationTests(unittest.TestCase):
    def setUp(self) -> None:
        build_docx.counters["equation"] = 0

    def _document_xml(self, latex: str) -> str:
        document = Document()
        build_docx.equation(document, latex)
        with tempfile.TemporaryDirectory(dir=Path(__file__).parents[1]) as tmp:
            path = Path(tmp) / "equation.docx"
            document.save(path)
            with zipfile.ZipFile(path) as archive:
                return archive.read("word/document.xml").decode("utf-8")

    def test_fraction_is_native_omml_and_numbered(self) -> None:
        xml = self._document_xml(
            r"\operatorname{CRF}(r,n)=\frac{r(1+r)^n}{(1+r)^n-1}"
        )

        self.assertIn("<m:oMath", xml)
        self.assertIn("<m:f>", xml)
        self.assertIn("(1)", xml)
        self.assertNotIn(r"\frac", xml)

    def test_piecewise_equation_is_native_omml(self) -> None:
        xml = self._document_xml(
            r"\bar p_t=\begin{cases}0,&v_t<v_{\mathrm{in}}\\1,&v_t\ge v_{\mathrm{r}}\end{cases}"
        )

        self.assertIn("<m:oMath", xml)
        self.assertIn("<m:m>", xml)
        self.assertNotIn(r"\begin{cases}", xml)

    def test_inline_math_is_native_omml(self) -> None:
        document = Document()
        build_docx.paragraph(
            document,
            r"The factor $\eta_l$ links input $f_l$ to output $q_l$.",
        )
        with tempfile.TemporaryDirectory(dir=Path(__file__).parents[1]) as tmp:
            path = Path(tmp) / "inline-equation.docx"
            document.save(path)
            with zipfile.ZipFile(path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")

        self.assertEqual(xml.count("<m:oMath"), 3)
        self.assertNotIn(r"$\eta_l$", xml)

    def test_markdown_emphasis_markers_do_not_leak_into_word(self) -> None:
        document = Document()
        build_docx.paragraph(
            document,
            "An *estimated allocation*, **result**, and `commit-id`.",
        )
        with tempfile.TemporaryDirectory(dir=Path(__file__).parents[1]) as tmp:
            path = Path(tmp) / "emphasis.docx"
            document.save(path)
            with zipfile.ZipFile(path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")

        self.assertNotIn("*estimated allocation*", xml)
        self.assertNotIn("**result**", xml)
        self.assertNotIn("`commit-id`", xml)

    def test_each_section_footer_has_one_page_field(self) -> None:
        document = Document()
        document.add_section()
        build_docx.add_page_numbers(document)
        with tempfile.TemporaryDirectory(dir=Path(__file__).parents[1]) as tmp:
            path = Path(tmp) / "page-numbers.docx"
            document.save(path)
            with zipfile.ZipFile(path) as archive:
                footers = [
                    archive.read(name).decode("utf-8")
                    for name in archive.namelist()
                    if name.startswith("word/footer") and name.endswith(".xml")
                ]

        self.assertEqual(len(footers), 2)
        self.assertTrue(all(footer.count('w:instr="PAGE"') == 1 for footer in footers))

    def test_word_is_asked_to_refresh_generated_fields(self) -> None:
        document = Document()
        build_docx.request_field_updates(document)
        with tempfile.TemporaryDirectory(dir=Path(__file__).parents[1]) as tmp:
            path = Path(tmp) / "field-refresh.docx"
            document.save(path)
            with zipfile.ZipFile(path) as archive:
                settings = archive.read("word/settings.xml").decode("utf-8")

        self.assertIn('<w:updateFields w:val="true"/>', settings)

    def test_headings_expose_outline_levels_to_word(self) -> None:
        document = Document()
        build_docx.heading(document, "Chapter", 20, toc_level=1)
        build_docx.heading(document, "Section", 16, toc_level=2)
        xml = document._element.xml

        self.assertIn('w:val="Heading1"', xml)
        self.assertIn('w:val="Heading2"', xml)
        self.assertGreaterEqual(xml.count('w:color w:val="000000"'), 2)

    def test_captions_use_sequence_fields_for_word_lists(self) -> None:
        document = Document()
        build_docx.caption(document, "Table 1: Parameters", above=True)
        build_docx.caption(document, "Figure 1: Network", above=False)
        xml = document._element.xml

        self.assertIn('w:instr="SEQ Table \\* ARABIC"', xml)
        self.assertIn('w:instr="SEQ Figure \\* ARABIC"', xml)


if __name__ == "__main__":
    unittest.main()
